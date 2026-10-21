"""Repeatable import of the PNE Tracker workbook into the SQLite curriculum catalog.

Reads only `07 Catálogo` and `06 Evidencias`. Curricular cells are stored
verbatim. The `Logrado ✓` column (student progress) is never stored.

Idempotent per baseline version:
  - same version + same workbook SHA-256  -> no changes ("unchanged")
  - same version + different SHA-256      -> refused (create a new version)

Every structural or consistency problem aborts the whole import (fail closed);
nothing is written unless all checks pass.

Usage:
    python -m curriculum.importer --workbook PATH [--db data/curriculum.sqlite]
        [--version PNE2026-SUM-v1.0] [--manifest references/MANIFEST.sha256]
        [--report report.json]
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
import sys
from collections import Counter, OrderedDict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

REPO_ROOT = Path(__file__).resolve().parents[1]
SCHEMA_PATH = Path(__file__).resolve().parent / "schema.sql"
DEFAULT_DB = REPO_ROOT / "data" / "curriculum.sqlite"
DEFAULT_MANIFEST = REPO_ROOT / "references" / "MANIFEST.sha256"
DEFAULT_VERSION = "PNE2026-SUM-v1.0"
SPEC = "PNE2026"
IMPORTER_VERSION = "1.0.0"

CATALOG_SHEET = "07 Catálogo"
EVIDENCE_SHEET = "06 Evidencias"
HEADER_ROW = 4
CATALOG_HEADERS = (
    "Código", "Materia", "Bloque MEP", "Numeración", "Afirmación / criterio",
    "Conceptos detallados", "Actividad de comprobación", "Fuente / página",
    "Cantidad evidencias", "Ítems previstos",
)
EVIDENCE_HEADERS = (
    "Objetivo", "N.º", "Evidencia / criterio verificable", "Materia",
    "Logrado ✓", "Afirmación asociada", "Bloque temático", "Fuente",
)
# Columns of 06 Evidencias that duplicate the catalogue: checked, never stored.
EVIDENCE_COPIES = {
    "Materia": "Materia",
    "Afirmación asociada": "Afirmación / criterio",
    "Bloque temático": "Bloque MEP",
    "Fuente": "Fuente / página",
}

SUBJECTS = {
    "Español": ("espanol", "ESP"),
    "Matemáticas": ("matematicas", "MAT"),
    "Ciencias": ("ciencias", "CIE"),
    "Estudios Sociales": ("estudios_sociales", "EST"),
}
# Workbook code prefix -> (subject name, assessment type, ID application code)
PREFIXES = {
    "EL": ("Español", "multiple_choice", "SUM"),
    "EE": ("Español", "written_rubric", "PE"),
    "M": ("Matemáticas", "multiple_choice", "SUM"),
    "C": ("Ciencias", "multiple_choice", "SUM"),
    "ES": ("Estudios Sociales", "multiple_choice", "SUM"),
}
CITED_DOCUMENTS = {
    "marco sumativo MEP 2026": ("MEP-MARCO-SUM-2026", "Marco sumativo MEP 2026", "spec_table"),
    "rúbrica MEP 2026": ("MEP-RUBRICA-PE-2026", "Rúbrica MEP 2026 (producción escrita)", "rubric"),
}
WORKBOOK_DOC_ID = "WORKBOOK-PNE-TRACKER-2026"

RE_CODE = re.compile(r"^([A-Z]+)(\d{2})$")
RE_NUMBERING = re.compile(r"^Afirmación (\d+)$")
RE_WEIGHT = re.compile(r"\((\d+) %\)$")
RE_PLANNED = re.compile(r"^(\d+) ítem\(s\) previsto\(s\)$")
RE_CITATION = re.compile(r"^(?:(Tabla \d+), )?(marco sumativo MEP 2026|rúbrica MEP 2026); pág\. PDF (\d+)\.$")


class WorkbookImportError(Exception):
    """Raised when the workbook fails validation; nothing is written."""

    def __init__(self, problems: list[str]):
        super().__init__(f"{len(problems)} problem(s): " + "; ".join(problems[:5]))
        self.problems = problems


@dataclass
class ParsedBaseline:
    blocks: list[dict[str, Any]] = field(default_factory=list)
    assertions: list[dict[str, Any]] = field(default_factory=list)
    evidences: list[dict[str, Any]] = field(default_factory=list)
    cited_documents: dict[str, tuple[str, str, str]] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- reading

def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def manifest_name(manifest: Path, sha: str) -> str:
    """Return the generic file name recorded for this checksum in the manifest."""
    for line in manifest.read_text(encoding="utf-8").splitlines():
        if line.startswith("#") or not line.strip():
            continue
        digest, _, name = line.partition(" ")
        if digest == sha:
            return name.lstrip("*").strip()
    raise WorkbookImportError([f"workbook SHA-256 {sha} is not registered in {manifest.name}"])


def read_sheet(workbook: Any, name: str, expected: tuple[str, ...]) -> list[tuple[int, dict[str, Any]]]:
    if name not in workbook.sheetnames:
        raise WorkbookImportError([f"sheet {name!r} not found"])
    rows = list(workbook[name].iter_rows(min_row=HEADER_ROW, values_only=True))
    header = tuple(rows[0][: len(expected)]) if rows else ()
    if header != expected:
        raise WorkbookImportError([f"{name}: header row {HEADER_ROW} is {header}, expected {expected}"])
    data = []
    for offset, values in enumerate(rows[1:], start=HEADER_ROW + 1):
        values = tuple(values[: len(expected)]) + (None,) * (len(expected) - len(values))
        if all(v is None for v in values):
            continue
        data.append((offset, dict(zip(expected, values))))
    return data


def as_int(value: Any, where: str, problems: list[str]) -> int | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
        problems.append(f"{where}: expected a whole number, got {value!r}")
        return None
    return int(value)


def as_text(value: Any, where: str, problems: list[str]) -> str:
    if not isinstance(value, str) or not value.strip():
        problems.append(f"{where}: missing text")
        return ""
    return value  # verbatim; never stripped or normalized


# --------------------------------------------------------------------------- parsing

def parse(workbook: Any) -> ParsedBaseline:
    problems: list[str] = []
    out = ParsedBaseline()
    catalog = read_sheet(workbook, CATALOG_SHEET, CATALOG_HEADERS)
    evidence_rows = read_sheet(workbook, EVIDENCE_SHEET, EVIDENCE_HEADERS)

    block_ids: "OrderedDict[tuple[str, str, str], str]" = OrderedDict()
    block_counter: Counter[tuple[str, str]] = Counter()
    by_code: dict[str, dict[str, Any]] = {}
    raw_by_code: dict[str, dict[str, Any]] = {}

    for row, cells in catalog:
        where = f"{CATALOG_SHEET}!row {row}"
        code = as_text(cells["Código"], f"{where} Código", problems)
        m = RE_CODE.match(code)
        if not m or m.group(1) not in PREFIXES:
            problems.append(f"{where}: unknown objective code {code!r}")
            continue
        if code in by_code:
            problems.append(f"{where}: duplicate objective code {code}")
            continue
        subject_name, assessment_type, app = PREFIXES[m.group(1)]
        if cells["Materia"] != subject_name:
            problems.append(f"{where}: code {code} implies {subject_name!r}, Materia is {cells['Materia']!r}")
            continue
        subject_code, subject_id = SUBJECTS[subject_name]

        planned_raw = as_text(cells["Ítems previstos"], f"{where} Ítems previstos", problems)
        planned_match = RE_PLANNED.match(planned_raw)
        if assessment_type == "written_rubric":
            if planned_raw != "criterio de rúbrica":
                problems.append(f"{where}: rubric code {code} has Ítems previstos {planned_raw!r}")
            planned = None
        elif planned_match:
            planned = int(planned_match.group(1))
        else:
            problems.append(f"{where}: unparseable Ítems previstos {planned_raw!r}")
            planned = None

        block_name = as_text(cells["Bloque MEP"], f"{where} Bloque MEP", problems)
        block_key = (subject_code, assessment_type, block_name)
        if block_key not in block_ids:
            block_counter[(subject_code, assessment_type)] += 1
            ordinal = block_counter[(subject_code, assessment_type)]
            block_id = f"{SPEC}-PRI-{subject_id}-{app}-B{ordinal}"
            weight = RE_WEIGHT.search(block_name)
            block_ids[block_key] = block_id
            out.blocks.append({
                "id": block_id, "subject_code": subject_code, "assessment_type": assessment_type,
                "ordinal": ordinal, "name_es": block_name,
                "weight_pct": int(weight.group(1)) if weight else None,
            })
        block_id = block_ids[block_key]

        numbering = as_text(cells["Numeración"], f"{where} Numeración", problems)
        num_match = RE_NUMBERING.match(numbering)
        if not num_match:
            problems.append(f"{where}: unparseable Numeración {numbering!r}")
            continue

        citation = as_text(cells["Fuente / página"], f"{where} Fuente / página", problems)
        cite = RE_CITATION.match(citation)
        if not cite:
            problems.append(f"{where}: unparseable source citation {citation!r}")
            continue
        doc_id, doc_title, doc_kind = CITED_DOCUMENTS[cite.group(2)]
        out.cited_documents[doc_id] = (doc_id, doc_title, doc_kind)

        evidence_count = as_int(cells["Cantidad evidencias"], f"{where} Cantidad evidencias", problems)
        record = {
            "id": f"{block_id}-A{int(num_match.group(1))}",
            "workbook_code": code,
            "block_id": block_id,
            "ordinal": int(num_match.group(1)),
            "numbering_es": numbering,
            "text_es": as_text(cells["Afirmación / criterio"], f"{where} Afirmación", problems),
            "concepts_es": as_text(cells["Conceptos detallados"], f"{where} Conceptos", problems),
            "check_activity_es": as_text(cells["Actividad de comprobación"], f"{where} Actividad", problems),
            "source_citation_es": citation,
            "source_document_id": doc_id,
            "source_table": cite.group(1),
            "source_page": int(cite.group(3)),
            "planned_items_raw": planned_raw,
            "planned_items": planned,
            "evidence_count": evidence_count,
        }
        by_code[code] = record
        raw_by_code[code] = cells
        out.assertions.append(record)

    duplicate_ids = [k for k, n in Counter(a["id"] for a in out.assertions).items() if n > 1]
    if duplicate_ids:
        problems.append(f"duplicate afirmación numbering within a bloque: {duplicate_ids}")

    ordinals_seen: dict[str, list[int]] = {}
    for row, cells in evidence_rows:
        where = f"{EVIDENCE_SHEET}!row {row}"
        code = cells["Objetivo"]
        parent = by_code.get(code)
        if parent is None:
            problems.append(f"{where}: Objetivo {code!r} not in {CATALOG_SHEET}")
            continue
        for column, catalog_column in EVIDENCE_COPIES.items():
            if cells[column] != raw_by_code[code][catalog_column]:
                problems.append(f"{where}: {column!r} differs from catalogue {catalog_column!r} for {code}")
        ordinal = as_int(cells["N.º"], f"{where} N.º", problems)
        if ordinal is None:
            continue
        ordinals_seen.setdefault(code, []).append(ordinal)
        out.evidences.append({
            "id": f"{parent['id']}-E{ordinal}",
            "assertion_id": parent["id"],
            "workbook_code": code,
            "ordinal": ordinal,
            "text_es": as_text(cells["Evidencia / criterio verificable"], f"{where} Evidencia", problems),
        })
        # cells["Logrado ✓"] is intentionally never read into the catalog (student data).

    for code, record in by_code.items():
        ordinals = ordinals_seen.get(code, [])
        if not ordinals:
            problems.append(f"{code}: no evidencias")
        elif sorted(ordinals) != list(range(1, len(ordinals) + 1)):
            problems.append(f"{code}: evidencia numbering is not 1..n: {ordinals}")
        if record["evidence_count"] is not None and record["evidence_count"] != len(ordinals):
            problems.append(f"{code}: Cantidad evidencias {record['evidence_count']} != {len(ordinals)} rows")

    duplicate_evidence = [k for k, n in Counter(e["id"] for e in out.evidences).items() if n > 1]
    if duplicate_evidence:
        problems.append(f"duplicate evidencia IDs: {duplicate_evidence[:5]}")

    texts = Counter((e["assertion_id"], e["text_es"]) for e in out.evidences)
    out.warnings += [f"repeated evidencia text in {aid}" for (aid, _), n in texts.items() if n > 1]
    similar = Counter(a["text_es"].strip().casefold() for a in out.assertions)
    out.warnings += [f"objectives share identical text: {t[:60]!r}" for t, n in similar.items() if n > 1]

    if problems:
        raise WorkbookImportError(problems)
    return out


# --------------------------------------------------------------------------- writing

def connect(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(SCHEMA_PATH.read_text(encoding="utf-8"))
    return conn


def _write(conn: sqlite3.Connection, version: str, parsed: ParsedBaseline, sha: str, file_ref: str) -> None:
    conn.executemany(
        "INSERT OR IGNORE INTO subject (code, name_es, id_code) VALUES (?, ?, ?)",
        [(code, name, id_code) for name, (code, id_code) in SUBJECTS.items()],
    )
    conn.execute(
        "INSERT INTO source_document VALUES (?, ?, ?, 'workbook', 2026, ?, ?, 'secondary', 'held', ?)",
        (version, WORKBOOK_DOC_ID, "PNE Tracker 2026 — Objetivos MEP (workbook)", file_ref, sha,
         datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")),
    )
    conn.executemany(
        "INSERT INTO source_document VALUES (?, ?, ?, ?, 2026, NULL, NULL, 'official', 'cited_not_held', NULL)",
        [(version, doc_id, title, kind) for doc_id, title, kind in parsed.cited_documents.values()],
    )
    conn.executemany(
        "INSERT INTO block VALUES (:v, :id, :subject_code, :assessment_type, :ordinal, :name_es, :weight_pct)",
        [dict(b, v=version) for b in parsed.blocks],
    )
    conn.executemany(
        """INSERT INTO assertion VALUES (:v, :id, :workbook_code, :block_id, :ordinal, :numbering_es,
           :text_es, :concepts_es, :check_activity_es, :source_citation_es, :source_document_id,
           :source_table, :source_page, :planned_items_raw, :planned_items, :evidence_count,
           'secondary', 'approved', 'unverified')""",
        [dict(a, v=version) for a in parsed.assertions],
    )
    conn.executemany(
        "INSERT INTO evidence VALUES (:v, :id, :assertion_id, :ordinal, :text_es, 'secondary', 'unverified')",
        [dict(e, v=version) for e in parsed.evidences],
    )


def summarize(conn: sqlite3.Connection, version: str) -> dict[str, Any]:
    components = conn.execute(
        """SELECT s.name_es AS subject, b.assessment_type, COUNT(DISTINCT a.id) AS objectives,
                  COUNT(e.id) AS evidences, COUNT(DISTINCT b.id) AS blocks
           FROM assertion a JOIN block b ON b.baseline_version = a.baseline_version AND b.id = a.block_id
           JOIN subject s ON s.code = b.subject_code
           LEFT JOIN evidence e ON e.baseline_version = a.baseline_version AND e.assertion_id = a.id
           WHERE a.baseline_version = ? GROUP BY s.name_es, b.assessment_type ORDER BY 1, 2""",
        (version,),
    ).fetchall()
    one = lambda sql: conn.execute(sql, (version,)).fetchone()[0]  # noqa: E731
    return {
        "baseline_version": version,
        "assertions": one("SELECT COUNT(*) FROM assertion WHERE baseline_version = ?"),
        "evidences": one("SELECT COUNT(*) FROM evidence WHERE baseline_version = ?"),
        "blocks": one("SELECT COUNT(*) FROM block WHERE baseline_version = ?"),
        "source_documents": one("SELECT COUNT(*) FROM source_document WHERE baseline_version = ?"),
        "components": [dict(r) for r in components],
    }


def import_workbook(workbook_path: Path, db_path: Path = DEFAULT_DB, version: str = DEFAULT_VERSION,
                    manifest: Path = DEFAULT_MANIFEST) -> dict[str, Any]:
    """Import the workbook as baseline `version`. Returns an import report dict."""
    import openpyxl  # local import: only the importer needs it

    sha = sha256_file(workbook_path)
    file_ref = manifest_name(manifest, sha)
    conn = connect(db_path)
    try:
        existing = conn.execute(
            "SELECT sha256 FROM source_document WHERE baseline_version = ? AND id = ?",
            (version, WORKBOOK_DOC_ID),
        ).fetchone()
        if existing and existing["sha256"] == sha:
            return {"status": "unchanged", "workbook_sha256": sha, "warnings": [], **summarize(conn, version)}
        if existing:
            raise WorkbookImportError([f"baseline {version} already imported from a different workbook "
                                f"({existing['sha256'][:12]}…); create a new baseline version"])

        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            parsed = parse(workbook)
        finally:
            workbook.close()
        with conn:  # single transaction: all or nothing
            _write(conn, version, parsed, sha, file_ref)
        return {"status": "imported", "workbook_sha256": sha, "importer_version": IMPORTER_VERSION,
                "warnings": parsed.warnings, **summarize(conn, version)}
    finally:
        conn.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Import the PNE Tracker workbook into the curriculum catalog.")
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--version", default=DEFAULT_VERSION)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--report", type=Path, help="write the JSON report here")
    args = parser.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        report = import_workbook(args.workbook, args.db, args.version, args.manifest)
    except WorkbookImportError as exc:
        print(json.dumps({"status": "failed", "problems": exc.problems}, ensure_ascii=False, indent=2))
        return 1
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        args.report.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
